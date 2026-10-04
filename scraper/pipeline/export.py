from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from scraper.log import get_logger

log = get_logger()

COLUMNS = [
    ("company", "Company"),
    ("title", "Title"),
    ("location", "Location"),
    ("department", "Department"),
    ("posted", "Posted"),
    ("employment_type", "Employment Type"),
    ("score", "Score"),
    ("ai_fit", "AI Fit"),
    ("ai_score", "AI Score"),
    ("ai_reason", "AI Reason"),
    ("job_id", "Job ID"),
    ("apply_url", "Apply URL"),
]
_LABEL_TO_KEY = {label: key for key, label in COLUMNS}


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(v) for v in value)
    elif isinstance(value, dict):
        value = str(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32_000]
    return value


def write_jobs(path: Path, jobs: list[dict], sheet_title: str = "Jobs") -> Path:
    present = {k for job in jobs for k in job}
    columns = [(k, label) for k, label in COLUMNS if k in present] or COLUMNS[:3]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([label for _, label in columns])
    for job in jobs:
        ws.append([_cell(job.get(k)) for k, _ in columns])

    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info(f"📊 Saved {len(jobs)} rows → {path}")
    return path


def read_jobs(path: Path) -> list[dict]:
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [
        _LABEL_TO_KEY.get(h, re.sub(r"\W+", "_", str(h or "").lower()).strip("_")) for h in next(rows, [])
    ]
    return [{k: v for k, v in zip(headers, row, strict=False) if v not in (None, "")} for row in rows]
